"""Text extraction (FR-2).

Handlers are registered per format. An unknown or unreadable format is never
rejected -- it falls back to filename/metadata indexing so the item is still
findable (FR-2, NFR-9).

Optional dependencies (pypdf, python-docx, openpyxl) are imported lazily so
that M1 runs on a bare interpreter and degrades to a clear status instead of
an import error.
"""

from __future__ import annotations

import csv
import io
import logging
import mimetypes
from dataclasses import dataclass, field
from pathlib import Path

log = logging.getLogger(__name__)

# Text kept per page/sheet so passages can carry page provenance.
@dataclass
class ExtractedText:
    text: str
    pages: list[str] = field(default_factory=list)
    status: str = "complete"          # complete | partial | failed
    note: str = ""                    # why, when not complete
    needs_ocr: bool = False

    @property
    def is_empty(self) -> bool:
        return not self.text.strip()


TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".log", ".json", ".yaml",
                 ".yml", ".ini", ".cfg", ".toml", ".html", ".htm", ".xml",
                 # Subtitles are plain text and are often the only searchable
                 # record of what was said in a video. 140 of them sat
                 # unindexed here purely because the suffix was missing.
                 ".srt", ".vtt", ".sub", ".ass"}
CSV_SUFFIXES = {".csv", ".tsv"}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
XLSX_SUFFIXES = {".xlsx", ".xlsm"}
# The old binary Excel format, which openpyxl cannot read. Banks also export
# HTML tables and CSV under this extension, which is why the handler sniffs the
# content rather than trusting the name.
XLS_SUFFIXES = {".xls"}
EML_SUFFIXES = {".eml"}
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".heic", ".heif", ".gif", ".tiff",
                  ".tif", ".webp", ".bmp"}

# Text by encoding, but not by intent: machine output with no sentence in it.
# These pass the "looks like text" sniff in _extract_unknown and then chunk into
# tens of thousands of passages each -- one 3D-printing .gcode file produced
# 104,227, more than four times the entire document corpus that surrounds it.
# Nobody searches for `G1 X92.7 Y104.5 E.03`, and the noise buries what they do
# search for. Indexed by filename and path, contents skipped.
MACHINE_SUFFIXES = {
    ".gcode", ".gco", ".g",          # 3D printer / CNC toolpaths
    ".nc", ".tap",                   # CNC
    ".stl", ".obj", ".3mf", ".amf",  # meshes (ASCII STL/OBJ sniff as text)
    ".ply", ".step", ".stp", ".iges", ".igs",
    ".map", ".sym", ".lst",          # linker and compiler output
    ".pack", ".idx",                 # VCS packfiles
    ".min.js", ".min.css",           # minified bundles
    ".ipynb_checkpoints",
    ".dump", ".mdump",
}


def guess_mime(path: Path) -> str:
    mime, _ = mimetypes.guess_type(str(path))
    if mime:
        return mime
    suffix = path.suffix.lower()
    if suffix in {".heic", ".heif"}:
        return "image/heic"
    if suffix == ".md":
        return "text/markdown"
    return "application/octet-stream"


def extract(path: Path, *, cfg=None) -> ExtractedText:
    """Extract text from a file. Never raises for content reasons.

    `cfg` is optional and only reaches the image path, where it enables the
    vision-model fallback for photographed documents Tesseract cannot read.
    Without it extraction behaves exactly as before.
    """
    suffix = path.suffix.lower()
    try:
        if suffix in PDF_SUFFIXES:
            return _extract_pdf(path)
        if suffix in CSV_SUFFIXES:
            return _extract_csv(path)
        if suffix in DOCX_SUFFIXES:
            return _extract_docx(path)
        if suffix in XLSX_SUFFIXES:
            return _extract_xlsx(path)
        if suffix in XLS_SUFFIXES:
            return _extract_xls(path)
        if suffix in EML_SUFFIXES:
            return _extract_eml(path)
        if suffix in IMAGE_SUFFIXES:
            return _extract_image(path, cfg=cfg)
        if suffix in TEXT_SUFFIXES:
            return _extract_plaintext(path)
        return _extract_unknown(path)
    except Exception as exc:  # extraction must not kill the pipeline
        log.warning("extraction failed for %s: %s", path, exc)
        return ExtractedText(text="", status="failed", note=f"{type(exc).__name__}: {exc}")


def _read_bytes(path: Path, limit: int | None = None) -> bytes:
    with open(path, "rb") as fh:
        return fh.read() if limit is None else fh.read(limit)


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _extract_plaintext(path: Path) -> ExtractedText:
    text = _decode(_read_bytes(path))
    # Markup is not content: `<div class="row">` in the index matches nothing
    # anyone would type, and it crowds out the words that do.
    if path.suffix.lower() in {".html", ".htm"}:
        text = _html_to_text(text)
    return ExtractedText(text=text, pages=[text])


def _html_to_text(markup: str) -> str:
    """Visible text from HTML. Tags in the index match nothing anyone types."""
    from html.parser import HTMLParser

    class Collector(HTMLParser):
        def __init__(self) -> None:
            super().__init__(convert_charrefs=True)
            self.parts: list[str] = []
            self._skip = 0

        def handle_starttag(self, tag: str, attrs: object) -> None:
            if tag in ("script", "style"):
                self._skip += 1

        def handle_endtag(self, tag: str) -> None:
            if tag in ("script", "style") and self._skip:
                self._skip -= 1
            elif tag in ("tr", "p", "div", "br", "li", "h1", "h2", "h3"):
                self.parts.append("\n")
            elif tag in ("td", "th"):
                self.parts.append("\t")

        def handle_data(self, data: str) -> None:
            if not self._skip and data.strip():
                self.parts.append(data.strip())

    collector = Collector()
    try:
        collector.feed(markup)
    except Exception:                                    # noqa: BLE001
        return markup
    lines = ("".join(collector.parts)).splitlines()
    return "\n".join(line.strip() for line in lines if line.strip())


def _extract_xls(path: Path) -> ExtractedText:
    """The old Excel format, and the things banks mislabel as it.

    A real .xls is a binary OLE file that openpyxl cannot open. But a large
    share of files with this extension are not Excel at all: banks commonly
    export an HTML table or a CSV and name it .xls, which is why the content
    decides here rather than the extension. Four bank statements sat unindexed
    under "no extractor" because of this.
    """
    sample = _read_bytes(path, 2048)

    # A genuine OLE2 document starts with this signature.
    if sample.startswith(b"\xd0\xcf\x11\xe0"):
        try:
            import xlrd
        except ImportError:
            return ExtractedText(
                text="", status="partial",
                note="old Excel format needs xlrd: indexed by filename only")
        try:
            book = xlrd.open_workbook(str(path))
            pages = []
            for sheet in book.sheets():
                rows = ["\t".join(str(c.value) for c in sheet.row(i))
                        for i in range(sheet.nrows)]
                pages.append("\n".join(rows))
            joined = "\n\n".join(pages)
            return ExtractedText(text=joined, pages=pages or [""],
                                 status="complete" if joined.strip() else "partial")
        except Exception as exc:                          # noqa: BLE001
            return ExtractedText(text="", status="partial",
                                 note=f"could not read as Excel: {exc}")

    text = _decode(_read_bytes(path))
    lowered = text[:4000].lower()
    if "<table" in lowered or "<html" in lowered:
        cleaned = _html_to_text(text)
        return ExtractedText(text=cleaned, pages=[cleaned], status="complete",
                             note="HTML table saved as .xls")
    if text.strip():
        return ExtractedText(text=text, pages=[text], status="complete",
                             note="plain text saved as .xls")
    return ExtractedText(text="", status="partial",
                         note="unreadable .xls: indexed by filename only")


def _extract_unknown(path: Path) -> ExtractedText:
    """Unknown format: index by name, and by content if it looks like text.

    FR-2 -- never reject a file.
    """
    # Machine output is text by encoding and meaningless by content, so the
    # printable-character sniff below would happily ingest all of it.
    name = path.name.lower()
    if (path.suffix.lower() in MACHINE_SUFFIXES
            or any(name.endswith(suffix) for suffix in MACHINE_SUFFIXES)):
        return ExtractedText(
            text="", status="partial",
            note="machine-generated format: indexed by filename only")

    sample = _read_bytes(path, 8192)
    if b"\x00" in sample:
        return ExtractedText(text="", status="partial",
                             note="binary format: indexed by filename only")
    text = _decode(_read_bytes(path))
    printable = sum(ch.isprintable() or ch.isspace() for ch in text[:4000])
    if text and printable / max(len(text[:4000]), 1) > 0.85:
        return ExtractedText(text=text, pages=[text], status="partial",
                             note="unrecognized format, treated as text")
    return ExtractedText(text="", status="partial",
                         note="unrecognized format: indexed by filename only")


def _extract_image(path: Path, *, cfg=None) -> ExtractedText:
    """Images: OCR the text (screenshots, photographed receipts).

    An unreadable image is not a failure -- it stays partial and keeps its
    filename in the index, and photo tagging (M8) adds EXIF and captions.
    """
    from . import ocr

    # Pass cfg only when there is one, so anything holding the older
    # single-argument signature (a stub, an out-of-tree caller) still works.
    result = ocr.ocr_image(path, cfg=cfg) if cfg is not None else ocr.ocr_image(path)
    if result.usable:
        return ExtractedText(text=result.text, pages=[result.text],
                             status="complete",
                             note=f"OCR via {result.backend} "
                                  f"(confidence {result.confidence:.2f})")

    if not ocr.available_backends():
        return ExtractedText(text="", status="partial",
                             note="no OCR backend installed", needs_ocr=True)
    return ExtractedText(text="", status="partial",
                         note="OCR found no readable text", needs_ocr=True)


def _extract_pdf(path: Path) -> ExtractedText:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ExtractedText(text="", status="partial",
                             note="pypdf not installed")

    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:
            log.debug("page extract failed in %s: %s", path, exc)
            pages.append("")

    text = "\n\n".join(pages)
    if text.strip():
        return ExtractedText(text=text, pages=pages)

    # No text layer: a scan. OCR page by page so citations keep a page number.
    from . import ocr

    if not ocr.available_backends():
        return ExtractedText(text="", pages=pages, status="partial",
                             note="scanned PDF; no OCR backend installed",
                             needs_ocr=True)

    results = ocr.ocr_pdf(path)
    ocr_pages = [r.text if r.usable else "" for r in results]
    ocr_text = "\n\n".join(ocr_pages)
    if not ocr_text.strip():
        return ExtractedText(text="", pages=pages, status="partial",
                             note="scanned PDF; OCR found no readable text",
                             needs_ocr=True)

    backend = next((r.backend for r in results if r.usable), "ocr")
    mean = sum(r.confidence for r in results) / max(len(results), 1)
    return ExtractedText(text=ocr_text, pages=ocr_pages, status="complete",
                         note=f"OCR via {backend} (confidence {mean:.2f})")


def _extract_csv(path: Path) -> ExtractedText:
    raw = _decode(_read_bytes(path))
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        pass

    rows = list(csv.reader(io.StringIO(raw), delimiter=delimiter))
    if not rows:
        return ExtractedText(text="")

    # Repeat the header on every row so a row remains meaningful in isolation
    # once passages are split -- a row divorced from its header is unsearchable.
    header = rows[0]
    lines = [" | ".join(header)]
    for row in rows[1:]:
        if not any(cell.strip() for cell in row):
            continue
        pairs = [f"{header[i]}: {cell}" if i < len(header) else cell
                 for i, cell in enumerate(row) if cell.strip()]
        lines.append("; ".join(pairs))

    text = "\n".join(lines)
    return ExtractedText(text=text, pages=[text])


def _extract_docx(path: Path) -> ExtractedText:
    try:
        import docx
    except ImportError:
        return ExtractedText(text="", status="partial",
                             note="python-docx not installed")

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts)
    return ExtractedText(text=text, pages=[text])


def _extract_xlsx(path: Path) -> ExtractedText:
    try:
        import openpyxl
    except ImportError:
        return ExtractedText(text="", status="partial",
                             note="openpyxl not installed")

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    pages: list[str] = []
    for sheet in workbook.worksheets:
        lines = [f"[sheet: {sheet.title}]"]
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
        if len(lines) > 1:
            pages.append("\n".join(lines))
    workbook.close()

    text = "\n\n".join(pages)
    return ExtractedText(text=text, pages=pages)


def _extract_eml(path: Path) -> ExtractedText:
    from email import policy
    from email.parser import BytesParser

    msg = BytesParser(policy=policy.default).parse(open(path, "rb"))
    headers = [f"{name}: {msg.get(name)}" for name in
               ("From", "To", "Cc", "Subject", "Date") if msg.get(name)]

    body = ""
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() == "text/plain":
                try:
                    body = part.get_content()
                    break
                except Exception:
                    continue
    else:
        try:
            body = msg.get_content()
        except Exception:
            body = ""

    text = "\n".join(headers) + "\n\n" + (body or "")
    return ExtractedText(text=text.strip(), pages=[text])
