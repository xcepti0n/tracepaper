"""Text extraction (FR-2).

Handlers are registered per format. An unknown or unreadable format is never
rejected -- it falls back to filename/metadata indexing so the item is still
findable (FR-2, NFR-9).

Optional dependencies (pypdf, python-docx, openpyxl) are imported lazily so
that M1 runs on a bare interpreter and degrades to a clear status instead of
an import error.
"""

from __future__ import annotations

import csv
import io
import logging
import mimetypes
from dataclasses import dataclass, field
from pathlib import Path

log = logging.getLogger(__name__)

# Text kept per page/sheet so passages can carry page provenance.
@dataclass
class ExtractedText:
    text: str
    pages: list[str] = field(default_factory=list)
    status: str = "complete"          # complete | partial | failed
    note: str = ""                    # why, when not complete
    needs_ocr: bool = False

    @property
    def is_empty(self) -> bool:
        return not self.text.strip()


TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".log", ".json", ".yaml",
                 ".yml", ".ini", ".cfg", ".toml", ".html", ".htm", ".xml"}
CSV_SUFFIXES = {".csv", ".tsv"}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
XLSX_SUFFIXES = {".xlsx", ".xlsm"}
EML_SUFFIXES = {".eml"}
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".heic", ".heif", ".gif", ".tiff",
                  ".tif", ".webp", ".bmp"}


def guess_mime(path: Path) -> str:
    mime, _ = mimetypes.guess_type(str(path))
    if mime:
        return mime
    suffix = path.suffix.lower()
    if suffix in {".heic", ".heif"}:
        return "image/heic"
    if suffix == ".md":
        return "text/markdown"
    return "application/octet-stream"


def extract(path: Path) -> ExtractedText:
    """Extract text from a file. Never raises for content reasons."""
    suffix = path.suffix.lower()
    try:
        if suffix in PDF_SUFFIXES:
            return _extract_pdf(path)
        if suffix in CSV_SUFFIXES:
            return _extract_csv(path)
        if suffix in DOCX_SUFFIXES:
            return _extract_docx(path)
        if suffix in XLSX_SUFFIXES:
            return _extract_xlsx(path)
        if suffix in EML_SUFFIXES:
            return _extract_eml(path)
        if suffix in IMAGE_SUFFIXES:
            return _extract_image(path)
        if suffix in TEXT_SUFFIXES:
            return _extract_plaintext(path)
        return _extract_unknown(path)
    except Exception as exc:  # extraction must not kill the pipeline
        log.warning("extraction failed for %s: %s", path, exc)
        return ExtractedText(text="", status="failed", note=f"{type(exc).__name__}: {exc}")


def _read_bytes(path: Path, limit: int | None = None) -> bytes:
    with open(path, "rb") as fh:
        return fh.read() if limit is None else fh.read(limit)


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _extract_plaintext(path: Path) -> ExtractedText:
    text = _decode(_read_bytes(path))
    return ExtractedText(text=text, pages=[text])


def _extract_unknown(path: Path) -> ExtractedText:
    """Unknown format: index by name, and by content if it looks like text.

    FR-2 -- never reject a file.
    """
    sample = _read_bytes(path, 8192)
    if b"\x00" in sample:
        return ExtractedText(text="", status="partial",
                             note="binary format: indexed by filename only")
    text = _decode(_read_bytes(path))
    printable = sum(ch.isprintable() or ch.isspace() for ch in text[:4000])
    if text and printable / max(len(text[:4000]), 1) > 0.85:
        return ExtractedText(text=text, pages=[text], status="partial",
                             note="unrecognized format, treated as text")
    return ExtractedText(text="", status="partial",
                         note="unrecognized format: indexed by filename only")


def _extract_image(path: Path) -> ExtractedText:
    """Images: OCR the text (screenshots, photographed receipts).

    An unreadable image is not a failure -- it stays partial and keeps its
    filename in the index, and photo tagging (M8) adds EXIF and captions.
    """
    from . import ocr

    result = ocr.ocr_image(path)
    if result.usable:
        return ExtractedText(text=result.text, pages=[result.text],
                             status="complete",
                             note=f"OCR via {result.backend} "
                                  f"(confidence {result.confidence:.2f})")

    if not ocr.available_backends():
        return ExtractedText(text="", status="partial",
                             note="no OCR backend installed", needs_ocr=True)
    return ExtractedText(text="", status="partial",
                         note="OCR found no readable text", needs_ocr=True)


def _extract_pdf(path: Path) -> ExtractedText:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ExtractedText(text="", status="partial",
                             note="pypdf not installed")

    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:
            log.debug("page extract failed in %s: %s", path, exc)
            pages.append("")

    text = "\n\n".join(pages)
    if text.strip():
        return ExtractedText(text=text, pages=pages)

    # No text layer: a scan. OCR page by page so citations keep a page number.
    from . import ocr

    if not ocr.available_backends():
        return ExtractedText(text="", pages=pages, status="partial",
                             note="scanned PDF; no OCR backend installed",
                             needs_ocr=True)

    results = ocr.ocr_pdf(path)
    ocr_pages = [r.text if r.usable else "" for r in results]
    ocr_text = "\n\n".join(ocr_pages)
    if not ocr_text.strip():
        return ExtractedText(text="", pages=pages, status="partial",
                             note="scanned PDF; OCR found no readable text",
                             needs_ocr=True)

    backend = next((r.backend for r in results if r.usable), "ocr")
    mean = sum(r.confidence for r in results) / max(len(results), 1)
    return ExtractedText(text=ocr_text, pages=ocr_pages, status="complete",
                         note=f"OCR via {backend} (confidence {mean:.2f})")


def _extract_csv(path: Path) -> ExtractedText:
    raw = _decode(_read_bytes(path))
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        pass

    rows = list(csv.reader(io.StringIO(raw), delimiter=delimiter))
    if not rows:
        return ExtractedText(text="")

    # Repeat the header on every row so a row remains meaningful in isolation
    # once passages are split -- a row divorced from its header is unsearchable.
    header = rows[0]
    lines = [" | ".join(header)]
    for row in rows[1:]:
        if not any(cell.strip() for cell in row):
            continue
        pairs = [f"{header[i]}: {cell}" if i < len(header) else cell
                 for i, cell in enumerate(row) if cell.strip()]
        lines.append("; ".join(pairs))

    text = "\n".join(lines)
    return ExtractedText(text=text, pages=[text])


def _extract_docx(path: Path) -> ExtractedText:
    try:
        import docx
    except ImportError:
        return ExtractedText(text="", status="partial",
                             note="python-docx not installed")

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts)
    return ExtractedText(text=text, pages=[text])


def _extract_xlsx(path: Path) -> ExtractedText:
    try:
        import openpyxl
    except ImportError:
        return ExtractedText(text="", status="partial",
                             note="openpyxl not installed")

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    pages: list[str] = []
    for sheet in workbook.worksheets:
        lines = [f"[sheet: {sheet.title}]"]
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
        if len(lines) > 1:
            pages.append("\n".join(lines))
    workbook.close()

    text = "\n\n".join(pages)
    return ExtractedText(text=text, pages=pages)


def _extract_eml(path: Path) -> ExtractedText:
    from email import policy
    from email.parser import BytesParser

    msg = BytesParser(policy=policy.default).parse(open(path, "rb"))
    headers = [f"{name}: {msg.get(name)}" for name in
               ("From", "To", "Cc", "Subject", "Date") if msg.get(name)]

    body = ""
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() == "text/plain":
                try:
                    body = part.get_content()
                    break
                except Exception:
                    continue
    else:
        try:
            body = msg.get_content()
        except Exception:
            body = ""

    text = "\n".join(headers) + "\n\n" + (body or "")
    return ExtractedText(text=text.strip(), pages=[text])
